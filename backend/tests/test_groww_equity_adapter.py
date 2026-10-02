"""Groww equity adapter tests against small synthetic .xlsx files generated per test."""

from dataclasses import fields
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.adapters.base import BrokerAdapter, StatementFormatError
from app.adapters.groww_equity import GrowwEquityAdapter
from app.domain.holdings import NormalizedHolding

TITLE = "Holdings statement for stocks as on 01-09-2026"
HEADER = [
    "Stock Name", "ISIN", "Quantity", "Average buy price",
    "Buy value", "Closing price", "Closing value", "Unrealised P&L",
]
# Synthetic rows. Closing price / P&L deliberately differ from the average buy price.
ROW_A = ["SAMPLE TEXTILES LIMITED", "INE000B01026", 12.0, 31.45, 377.4, 9.5, 114.0, -263.4]
ROW_B = ["SYNTHETIC POWER LIMITED", "INE000A01011", 5, 1234.5678, 6172.84, 999.0, 4995.0, -1177.84]
BATCH_ID = "batch-123"

adapter = GrowwEquityAdapter()


def statement_rows(data, *, header=HEADER, title=TITLE, client_code=True, extra_rows_before_header=()):
    """The real Groww layout: identity block, title, summary block, blank row, header, data."""
    return [
        ["Name", "Synthetic Holder"],
        ["Unique Client Code", "TEST0001"] if client_code else [],
        [],
        [title] if title else [],
        [],
        ["Summary"],
        ["Invested Value", 1000.0],
        ["Closing Value", 913.75],
        ["Unrealised P&L", -86.25],
        *extra_rows_before_header,
        [],
        header,
        *data,
    ]


@pytest.fixture
def write_xlsx(tmp_path):
    def write(sheets: dict[str, list[list]], *, column_offset: int = 0, name: str = "statement.xlsx") -> str:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row_number, row in enumerate(rows, start=1):
                for column_number, value in enumerate(row, start=1 + column_offset):
                    if value is not None:
                        sheet.cell(row_number, column_number, value)
        path = tmp_path / name
        workbook.save(path)
        return str(path)

    return write


@pytest.fixture
def statement(write_xlsx):
    """Write a standard single-sheet Groww statement with the given data rows."""
    return lambda data, **layout: write_xlsx({"Sheet1": statement_rows(data, **layout)})


HEADER_ROW = 11  # in statement_rows' default layout (as in the real file); data starts on row 12


# --- Interface and recognition -------------------------------------------------------------

def test_implements_broker_adapter_interface():
    assert isinstance(adapter, BrokerAdapter)


def test_can_parse_recognises_groww_statement(statement):
    assert adapter.can_parse(statement([ROW_A])) is True


def test_can_parse_rejects_other_files(write_xlsx, tmp_path):
    zerodha_rows = [["Client ID", "AB1234"], ["Equity Holdings Statement as on 2026-09-02"]]
    zerodha = write_xlsx({"Combined": zerodha_rows}, name="z.xlsx")
    mf_header = ["Scheme Name", "Transaction Type", "Units", "NAV", "Amount", "Date"]
    groww_mf = write_xlsx({"Transactions": [mf_header]}, name="m.xlsx")
    no_client_code = write_xlsx({"Sheet1": statement_rows([ROW_A], client_code=False)}, name="n.xlsx")
    not_xlsx = tmp_path / "notes.xlsx"
    not_xlsx.write_text("not a spreadsheet")
    for path in (zerodha, groww_mf, no_client_code, str(not_xlsx), str(tmp_path / "missing.xlsx")):
        assert adapter.can_parse(path) is False


def test_parse_rejects_unreadable_or_foreign_files(write_xlsx, tmp_path):
    not_xlsx = tmp_path / "notes.xlsx"
    not_xlsx.write_text("not a spreadsheet")
    with pytest.raises(StatementFormatError, match="readable"):
        adapter.parse(str(not_xlsx), BATCH_ID)
    other = write_xlsx({"Sheet1": [["Something else"]]})
    with pytest.raises(StatementFormatError, match="not a Groww stocks holdings statement"):
        adapter.parse(other, BATCH_ID)


# --- Valid parsing and mapping ---------------------------------------------------------------

def test_parses_and_maps_every_valid_row(statement):
    result = adapter.parse(statement([ROW_A, ROW_B]), BATCH_ID)

    assert result.errors == []
    assert result.statement_date == date(2026, 9, 1)
    assert result.holdings == [
        NormalizedHolding(
            instrument_name="SAMPLE TEXTILES LIMITED", isin="INE000B01026", symbol=None, asset_class="equity",
            quantity=Decimal("12.0"), average_buy_price=Decimal("31.45"),
            native_currency="INR", source_broker="groww", import_batch_id=BATCH_ID,
        ),
        NormalizedHolding(
            instrument_name="SYNTHETIC POWER LIMITED", isin="INE000A01011", symbol=None, asset_class="equity",
            quantity=Decimal("5"), average_buy_price=Decimal("1234.5678"),
            native_currency="INR", source_broker="groww", import_batch_id=BATCH_ID,
        ),
    ]


def test_financial_values_are_exact_decimals(statement):
    holding = adapter.parse(statement([ROW_A]), BATCH_ID).holdings[0]
    assert type(holding.quantity) is Decimal and type(holding.average_buy_price) is Decimal
    assert str(holding.average_buy_price) == "31.45"  # not 31.449999999999999289457264239899814128875732421875


def test_closing_price_and_statement_pnl_are_not_holding_values(statement):
    holding = adapter.parse(statement([ROW_A]), BATCH_ID).holdings[0]
    assert holding.average_buy_price == Decimal("31.45")  # 'Average buy price', not 'Closing price' (9.5)
    assert {f.name for f in fields(NormalizedHolding)} == {
        "instrument_name", "isin", "symbol", "asset_class", "quantity",
        "average_buy_price", "native_currency", "source_broker", "import_batch_id",
    }


def test_symbol_is_not_invented_from_stock_name(statement):
    assert adapter.parse(statement([ROW_A]), BATCH_ID).holdings[0].symbol is None


def test_numeric_text_with_thousands_separators(statement):
    row = ["SYNTHETIC LTD", "INE000A01011", "1,000", "12,34,567.50", None, None, None, None]
    holding = adapter.parse(statement([row]), BATCH_ID).holdings[0]
    assert holding.quantity == Decimal("1000")
    assert holding.average_buy_price == Decimal("1234567.50")


def test_zero_average_buy_price_is_allowed(statement):
    row = ["BONUS SHARES LTD", "INE000A01011", 10, 0, 0, 5.0, 50.0, 50.0]
    result = adapter.parse(statement([row]), BATCH_ID)
    assert result.errors == []
    assert result.holdings[0].average_buy_price == Decimal("0")


def test_text_fields_are_trimmed(statement):
    row = ["  SAMPLE TEXTILES LIMITED ", " INE000B01026 ", 12.0, 31.45, None, None, None, None]
    holding = adapter.parse(statement([row]), BATCH_ID).holdings[0]
    assert (holding.instrument_name, holding.isin) == ("SAMPLE TEXTILES LIMITED", "INE000B01026")


# --- Sheet selection and header detection ----------------------------------------------------

def test_selects_statement_sheet_by_content(write_xlsx):
    decoy = [HEADER, ["DECOY LTD", "INE111A01011", 1, 1.0]]  # has the header but no Groww markers
    path = write_xlsx({"Notes": decoy, "Holdings": statement_rows([ROW_A])})
    result = adapter.parse(path, BATCH_ID)
    assert [h.instrument_name for h in result.holdings] == ["SAMPLE TEXTILES LIMITED"]


def test_header_found_at_shifted_row_and_column(write_xlsx):
    extra = [["Some note"], [], ["Another note"], []]
    rows = statement_rows([ROW_A, ["BAD LTD", "INE000B01026", "x", 1.0]], extra_rows_before_header=extra)
    result = adapter.parse(write_xlsx({"Sheet1": rows}, column_offset=3), BATCH_ID)
    assert len(result.holdings) == 1
    assert [e.row_number for e in result.errors] == [HEADER_ROW + len(extra) + 2]  # real Excel row number


def test_metadata_reusing_header_words_is_not_mistaken_for_header(statement):
    # The summary block's 'Unrealised P&L' (always present) and a lone 'Stock Name' label must not match.
    decoys = [["Stock Name", "see below"], ["ISIN", "Quantity"]]
    result = adapter.parse(statement([ROW_A], extra_rows_before_header=decoys), BATCH_ID)
    assert result.errors == []
    assert [h.isin for h in result.holdings] == ["INE000B01026"]


def test_columns_are_matched_by_name_not_position(statement):
    header = ["Average buy price", "Closing price", "Quantity", "ISIN", "Stock Name"]
    row = [31.45, 9.5, 12.0, "INE000B01026", "SAMPLE TEXTILES LIMITED"]
    holding = adapter.parse(statement([row], header=header), BATCH_ID).holdings[0]
    assert (holding.quantity, holding.average_buy_price) == (Decimal("12.0"), Decimal("31.45"))


@pytest.mark.parametrize("missing", ["Stock Name", "ISIN", "Quantity", "Average buy price"])
def test_missing_required_header_rejects_file(statement, missing):
    header = [h for h in HEADER if h != missing]
    with pytest.raises(StatementFormatError, match="header"):
        adapter.parse(statement([ROW_A[: len(header)]], header=header), BATCH_ID)


@pytest.mark.parametrize(
    "title",
    [
        "Holdings statement for stocks",
        "Holdings statement for stocks as on 31-02-2026",
        "Holdings statement for stocks as on 2026-09-01",
    ],
)
def test_missing_or_invalid_statement_date_rejects_file(statement, title):
    with pytest.raises(StatementFormatError, match="date|Holdings statement"):
        adapter.parse(statement([ROW_A], title=title), BATCH_ID)


# --- Blank rows ------------------------------------------------------------------------------

def test_blank_rows_are_skipped_not_reported(statement):
    whitespace_only = ["   ", None, " ", None]
    result = adapter.parse(statement([[], ROW_A, whitespace_only, [], ROW_B, [], []]), BATCH_ID)
    assert len(result.holdings) == 2
    assert result.errors == []


def test_statement_with_no_data_rows(statement):
    result = adapter.parse(statement([]), BATCH_ID)
    assert (result.holdings, result.errors) == ([], [])


# --- Row-level validation --------------------------------------------------------------------

@pytest.mark.parametrize(
    ("column", "bad_value", "field", "message"),
    [
        (0, None, "Stock Name", "Stock Name is missing"),
        (0, "   ", "Stock Name", "Stock Name is missing"),
        (1, None, "ISIN", "ISIN is missing"),
        (1, "INE000B0102", "ISIN", "ISIN is not a valid 12-character ISIN"),
        (1, "ine000b01026", "ISIN", "ISIN is not a valid 12-character ISIN"),
        (1, "INE000B0102X", "ISIN", "ISIN is not a valid 12-character ISIN"),
        (2, None, "Quantity", "Quantity is missing or not a number"),
        (2, "abc", "Quantity", "Quantity is not a number"),
        (2, "1,2", "Quantity", "Quantity is not a number"),
        (2, "1e5", "Quantity", "Quantity is not a number"),
        (2, True, "Quantity", "Quantity is missing or not a number"),
        (2, 0, "Quantity", "Quantity must be greater than zero"),
        (2, -5.0, "Quantity", "Quantity must be greater than zero"),
        (3, None, "Average buy price", "Average buy price is missing or not a number"),
        (3, "N/A", "Average buy price", "Average buy price is not a number"),
        (3, -1.5, "Average buy price", "Average buy price must not be negative"),
    ],
)
def test_invalid_value_reported_for_its_row_and_other_rows_still_parse(statement, column, bad_value, field, message):
    bad_row = list(ROW_B)
    bad_row[column] = bad_value
    result = adapter.parse(statement([ROW_A, bad_row, ROW_A]), BATCH_ID)

    assert len(result.holdings) == 2  # per-row results, not all-or-nothing
    assert len(result.errors) == 1
    error = result.errors[0]
    assert (error.row_number, error.field, error.message) == (HEADER_ROW + 2, field, message)
    assert result.failed_row_count == 1


def test_every_invalid_field_in_a_row_is_reported(statement):
    result = adapter.parse(statement([[None, "BAD", "abc", -1]]), BATCH_ID)
    assert result.holdings == []
    assert [e.field for e in result.errors] == ["Stock Name", "ISIN", "Quantity", "Average buy price"]
    assert result.failed_row_count == 1


def test_error_messages_never_echo_cell_values(statement):
    secret = "ACCOUNT-SECRET-42"
    result = adapter.parse(statement([[secret, secret, secret, secret]]), BATCH_ID)
    assert result.errors and all(secret not in e.message for e in result.errors)
