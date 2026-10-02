"""Shared synthetic fixtures for Module 7 tests (no real broker data)."""

import uuid
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from app.domain.holdings import NormalizedHolding
from app.models import User
from app.repositories import import_repository, user_repository
from tests import test_groww_equity_adapter as groww_equity
from tests import test_groww_mutual_fund_adapter as groww_mf
from tests import test_zerodha_equity_adapter as zerodha


def make_user(db, email: str, reporting_currency: str = "INR") -> User:
    user = user_repository.create(db, email, "not-a-real-hash")
    user.reporting_currency = reporting_currency
    db.commit()
    return user


def nh(name="SYNTHA", isin="INE000D01019", symbol="SYNTHA", qty="10", avg="100", *, asset_class="equity",
       broker="zerodha", currency="INR") -> NormalizedHolding:
    return NormalizedHolding(name, isin, symbol, asset_class, Decimal(qty), Decimal(avg), currency, broker, "unused")


def seed(db, user_id: int, broker: str, asset_class: str, statement_date: date, holdings) -> str:
    batch_id = str(uuid.uuid4())
    import_repository.create_batch(
        db, batch_id=batch_id, user_id=user_id, broker=broker, asset_class=asset_class,
        statement_date=statement_date, file_name=f"{broker}-{asset_class}.xlsx",
        rows_succeeded=len(holdings), rows_failed=0, holdings=holdings,
    )
    db.commit()
    return batch_id


def write_xlsx(path: Path, sheets: dict[str, list[list]], column_offset: int = 0) -> str:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1 + column_offset):
                if value is not None:
                    sheet.cell(r, c, value)
    workbook.save(path)
    return str(path)


def groww_equity_file(path: Path, rows=(groww_equity.ROW_A,), as_on="01-09-2026") -> str:
    title = f"Holdings statement for stocks as on {as_on}"
    return write_xlsx(path, {"Sheet1": groww_equity.statement_rows(list(rows), title=title)})


def zerodha_file(path: Path, holdings=(zerodha.HOLDING_X,), as_on="2026-06-30") -> str:
    title = f"Combined Holdings Statement as on {as_on}"
    return write_xlsx(path, {"Combined": zerodha.sheet_rows(list(holdings), title=title)}, column_offset=1)


def groww_mf_file(path: Path, transactions=None, to="Jun 30 2026") -> str:
    transactions = transactions if transactions is not None else [groww_mf.tx()]
    title = f"TRANSACTIONS FROM Jan 01 2025 TO {to}"
    return write_xlsx(path, {"Transactions": groww_mf.statement_rows(transactions, title=title)})
