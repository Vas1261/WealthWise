"""Groww mutual fund adapter tests (locked B3 rule) against small synthetic .xlsx files."""

from dataclasses import astuple
from datetime import date
from decimal import Context, Decimal

import openpyxl
import pytest

from app.adapters.base import BrokerAdapter, StatementFormatError
from app.adapters.groww_equity import GrowwEquityAdapter
from app.adapters.groww_mutual_fund import GrowwMutualFundAdapter
from app.adapters.zerodha_equity import ZerodhaEquityAdapter
from app.domain.holdings import NormalizedHolding

BATCH_ID = "batch-789"
TITLE = "TRANSACTIONS FROM Jan 01 2025 TO Jun 30 2026"
HEADER = ["Scheme Name", "Transaction Type", "Units", "NAV", "Amount", "Date"]
ALPHA = "Synthetic Alpha Fund Direct Growth"
BETA = "Synthetic Beta Index Fund Direct Growth"
# Invented personal details, used only to prove they never reach the output.
HOLDER, MOBILE, PAN = "Synthetic Holder", "0000000000", "AAAAA0000A"
B3 = Context(prec=28)  # the adapter's documented precision
UNSUPPORTED_TYPE = "Transaction Type is not supported; only PURCHASE and REDEEM are imported"


def tx(scheme=ALPHA, kind="PURCHASE", units="10", nav="20", amount="200", when="15 Jan 2026") -> dict:
    return {"Scheme Name": scheme, "Transaction Type": kind, "Units": units, "NAV": nav, "Amount": amount, "Date": when}


def statement_rows(transactions, *, header=HEADER, title=TITLE, extra_rows_before_header=()):
    """The real Groww layout: personal details rows 3-6, title row 10, header row 12."""
    return [
        [], [],
        ["Personal Details"],
        ["Name", HOLDER],
        ["Mobile Number", MOBILE],
        ["PAN", PAN],
        [], [], [],
        [title] if title else [],
        *extra_rows_before_header,
        [],
        header,
        *([t.get(h) for h in header] for t in transactions),
    ]


HEADER_ROW = 12  # in statement_rows' default layout (as in the real file); data starts on row 13


@pytest.fixture
def write_xlsx(tmp_path):
    def write(sheets: dict[str, list[list]], *, column_offset: int = 0, name: str = "orders.xlsx") -> str:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row_number, row in enumerate(rows, start=1):
                for column_number, value in enumerate(row, start=1 + column_offset):
                    if value is not None:
                        sheet.cell(row_number, column_number, value)
        path = tmp_path / name
        workbook.save(path)
        return str(path)

    return write


@pytest.fixture
def orders(write_xlsx):
    """Write a standard single-sheet Groww order history with the given transactions."""
    return lambda transactions, **layout: write_xlsx({"Transactions": statement_rows(transactions, **layout)})


adapter = GrowwMutualFundAdapter()


def parse(path):
    return adapter.parse(path, BATCH_ID)


def only_holding(result) -> NormalizedHolding:
    assert len(result.holdings) == 1
    return result.holdings[0]


# --- Interface, recognition, statement date -----------------------------------------------------

def test_implements_broker_adapter_interface():
    assert isinstance(adapter, BrokerAdapter)


def test_can_parse_recognises_order_history(orders):
    assert adapter.can_parse(orders([tx()])) is True


def test_statement_date_is_the_title_to_date_not_file_name(write_xlsx):
    name = "Mutual_Funds_Order_History_2019-04-01_2027-03-31.xlsx"
    path = write_xlsx({"Transactions": statement_rows([tx()])}, name=name)
    assert parse(path).statement_date == date(2026, 6, 30)


def test_aa_groww_equity_workbook_not_claimed(write_xlsx):
    groww_equity = write_xlsx({"Sheet1": [
        ["Name", HOLDER], ["Unique Client Code", "TEST0001"], [],
        ["Holdings statement for stocks as on 01-09-2026"], [],
        ["Stock Name", "ISIN", "Quantity", "Average buy price"],
        ["SYNTHETIC LTD", "INE000B01026", 1.0, 10.0],
    ]})
    assert GrowwEquityAdapter().can_parse(groww_equity) is True  # sanity: it is a valid Groww equity file
    assert adapter.can_parse(groww_equity) is False


def test_ab_zerodha_equity_workbook_not_claimed(write_xlsx):
    zerodha = write_xlsx({"Combined": [
        ["Client ID", "ZZ0000"], ["Combined Holdings Statement as on 2026-06-30"],
        ["Symbol", "ISIN", "Instrument Type", "Quantity Available", "Average Price"],
        ["SYNTHA", "INE000D01019", "-", 1.0, 10.0],
    ]}, column_offset=1)
    assert ZerodhaEquityAdapter().can_parse(zerodha) is True  # sanity: it is a valid Zerodha file
    assert adapter.can_parse(zerodha) is False


def test_ac_unrelated_files_not_claimed(write_xlsx, tmp_path):
    unrelated = write_xlsx({"Sheet1": [["Hello"], ["World", 1, 2]]}, name="u.xlsx")
    header_without_title = write_xlsx({"Transactions": statement_rows([tx()], title=None)}, name="h.xlsx")
    not_xlsx = tmp_path / "notes.xlsx"
    not_xlsx.write_text("not a spreadsheet")
    for path in (unrelated, header_without_title, str(not_xlsx), str(tmp_path / "missing.xlsx")):
        assert adapter.can_parse(path) is False


def test_equity_adapters_do_not_claim_order_history(orders):
    path = orders([tx()])
    assert GrowwEquityAdapter().can_parse(path) is False
    assert ZerodhaEquityAdapter().can_parse(path) is False


def test_parse_rejects_unreadable_and_foreign_files(write_xlsx, tmp_path):
    not_xlsx = tmp_path / "notes.xlsx"
    not_xlsx.write_text("not a spreadsheet")
    with pytest.raises(StatementFormatError, match="readable"):
        parse(str(not_xlsx))
    with pytest.raises(StatementFormatError, match="not a Groww mutual fund order history"):
        parse(write_xlsx({"Sheet1": [["Something else"]]}))


@pytest.mark.parametrize(
    "title",
    [
        "TRANSACTIONS FROM Jan 01 2025 TO Feb 30 2026",  # impossible date
        "TRANSACTIONS FROM Jul 01 2026 TO Jun 30 2026",  # FROM after TO
        "TRANSACTIONS FROM 2025-01-01 TO 2026-06-30",  # wrong date format
        "Transactions from Jan 01 2025 to Jun 30 2026",  # wrong case
        "TRANSACTIONS FROM Jan 01 2025",  # no TO date
        "TRANSACTIONS FROM Jan 1 2025 TO Jun 30 2026",  # single-digit day
    ],
)
def test_invalid_title_means_no_qualifying_sheet(orders, title):
    path = orders([tx()], title=title)
    assert adapter.can_parse(path) is False
    with pytest.raises(StatementFormatError):
        parse(path)


def test_title_whitespace_is_collapsed(orders):
    path = orders([tx()], title="  TRANSACTIONS  FROM Jan 01 2025\tTO Jun 30 2026 ")
    assert parse(path).statement_date == date(2026, 6, 30)


def test_duplicated_title_means_no_qualifying_sheet(orders):
    path = orders([tx()], extra_rows_before_header=[[TITLE]])
    assert adapter.can_parse(path) is False


def test_header_above_title_does_not_count(write_xlsx):
    rows = [HEADER, [], [TITLE], [], ["not the header"], *([tx()[h] for h in HEADER],)]
    assert adapter.can_parse(write_xlsx({"Transactions": rows})) is False


# --- Sheet selection ---------------------------------------------------------------------------

def test_qualifying_sheet_found_by_content_not_name(write_xlsx):
    path = write_xlsx({"Cover": [["Report"]], "Orders": statement_rows([tx()])})
    assert len(parse(path).holdings) == 1


def test_transactions_name_breaks_a_tie(write_xlsx):
    path = write_xlsx({
        "Copy": statement_rows([tx(scheme=BETA)]),
        "Transactions": statement_rows([tx(scheme=ALPHA)]),
    })
    assert [h.instrument_name for h in parse(path).holdings] == [ALPHA]


def test_two_qualifying_sheets_without_a_tie_breaker_are_ambiguous(write_xlsx):
    path = write_xlsx({"Orders A": statement_rows([tx()]), "Orders B": statement_rows([tx()])})
    assert adapter.can_parse(path) is False
    with pytest.raises(StatementFormatError, match="ambiguous"):
        parse(path)


def test_non_qualifying_sheet_does_not_block_qualifying_one(write_xlsx):
    path = write_xlsx({
        "Transactions": statement_rows([tx(scheme=BETA)], title="TRANSACTIONS FROM Jan 01 2025 TO Feb 30 2026"),
        "Orders": statement_rows([tx(scheme=ALPHA)]),
    })
    assert [h.instrument_name for h in parse(path).holdings] == [ALPHA]


# --- A-H, AE, AF, AG: B3 aggregation -------------------------------------------------------------

def test_a_single_purchase_maps_to_normalized_holding(orders):
    result = parse(orders([tx(units="7.375", nav="41.625")]))
    assert result.errors == []
    assert result.holdings == [
        NormalizedHolding(
            instrument_name=ALPHA, isin=None, symbol=None, asset_class="mutual_fund",
            quantity=Decimal("7.375"), average_buy_price=Decimal("41.625"),
            native_currency="INR", source_broker="groww", import_batch_id=BATCH_ID,
        )
    ]


def test_b_multiple_purchases_weighted_by_units(orders):
    h = only_holding(parse(orders([tx(units="100", nav="20"), tx(units="50", nav="24")])))
    assert h.quantity == Decimal("150")
    assert h.average_buy_price == B3.divide(Decimal(3200), Decimal(150))  # 21.333...
    assert str(h.average_buy_price).startswith("21.33333333333")


def test_c_purchase_and_redeem_specification_example(orders):
    rows = [tx(units="100", nav="20"), tx(units="50", nav="24"), tx(kind="REDEEM", units="30", nav="25")]
    h = only_holding(parse(orders(rows)))
    assert h.quantity == Decimal("120")
    assert h.average_buy_price == B3.divide(Decimal(3200), Decimal(150))


def test_redemption_nav_never_enters_the_average(orders):
    """Fails if REDEEM NAV were included: a units-weighted average over all rows would give 32, not 15."""
    rows = [tx(units="100", nav="10"), tx(units="100", nav="20"), tx(kind="REDEEM", units="50", nav="100")]
    h = only_holding(parse(orders(rows)))
    assert h.average_buy_price == Decimal("15")
    assert h.average_buy_price != Decimal("32")
    assert h.quantity == Decimal("150")


def test_d_multiple_purchases_and_redemptions(orders):
    rows = [
        tx(units="40", nav="10"), tx(kind="REDEEM", units="10", nav="12"),
        tx(units="60", nav="15"), tx(kind="REDEEM", units="25", nav="18"), tx(units="20", nav="30"),
    ]
    h = only_holding(parse(orders(rows)))
    assert h.quantity == Decimal("85")  # 40 + 60 + 20 - 10 - 25
    assert h.average_buy_price == B3.divide(Decimal(400 + 900 + 600), Decimal(120))


def test_e_redemption_order_does_not_change_average(orders):
    before = [tx(kind="REDEEM", units="5", nav="99"), tx(units="10", nav="20"), tx(units="10", nav="30")]
    after = [tx(units="10", nav="20"), tx(units="10", nav="30"), tx(kind="REDEEM", units="5", nav="1")]
    assert only_holding(parse(orders(before))).average_buy_price == Decimal("25")
    assert astuple(only_holding(parse(orders(before)))) == astuple(only_holding(parse(orders(after))))


def test_f_fully_redeemed_scheme_excluded(orders):
    rows = [tx(units="12.5"), tx(kind="REDEEM", units="12.5", nav="30"), tx(scheme=BETA)]
    result = parse(orders(rows))
    assert [h.instrument_name for h in result.holdings] == [BETA]
    assert result.errors == []  # exclusion is not an error


def test_g_over_redeemed_scheme_excluded(orders):
    result = parse(orders([tx(units="10"), tx(kind="REDEEM", units="10.001"), tx(scheme=BETA)]))
    assert [h.instrument_name for h in result.holdings] == [BETA]


def test_ae_redeem_only_scheme_excluded(orders):
    result = parse(orders([tx(scheme=ALPHA, kind="REDEEM", units="5"), tx(scheme=BETA)]))
    assert [h.instrument_name for h in result.holdings] == [BETA]
    assert result.errors == []


def test_h_schemes_kept_separate_in_first_appearance_order(orders):
    rows = [tx(scheme=BETA, units="2", nav="100"), tx(scheme=ALPHA, units="10", nav="20"),
            tx(scheme=BETA, units="2", nav="200"), tx(scheme=ALPHA, kind="REDEEM", units="4", nav="50")]
    result = parse(orders(rows))
    assert [(h.instrument_name, h.quantity, h.average_buy_price) for h in result.holdings] == [
        (BETA, Decimal("4"), Decimal("150")),
        (ALPHA, Decimal("6"), Decimal("20")),
    ]


def test_i_scheme_name_whitespace_collapsed_for_grouping(orders):
    rows = [tx(scheme=f"  {ALPHA} "), tx(scheme=ALPHA.replace(" ", "   ")), tx(scheme=ALPHA)]
    h = only_holding(parse(orders(rows)))
    assert (h.instrument_name, h.quantity) == (ALPHA, Decimal("30"))


def test_scheme_names_are_case_sensitive(orders):
    assert len(parse(orders([tx(scheme=ALPHA), tx(scheme=ALPHA.upper())])).holdings) == 2


def test_af_average_weights_purchase_units_only(orders):
    # Redeeming most units must not re-weight the average towards later purchases.
    rows = [tx(units="90", nav="10"), tx(units="10", nav="110"), tx(kind="REDEEM", units="95", nav="200")]
    h = only_holding(parse(orders(rows)))
    assert (h.quantity, h.average_buy_price) == (Decimal("5"), Decimal("20"))  # (900 + 1100) / 100


def test_ag_amount_never_affects_average(orders):
    plausible = [tx(units="10", nav="20", amount="200"), tx(units="10", nav="30", amount="300")]
    nonsense = [tx(units="10", nav="20", amount="999,999"), tx(units="10", nav="30", amount="1")]
    one, two = only_holding(parse(orders(plausible))), only_holding(parse(orders(nonsense)))
    assert one.average_buy_price == two.average_buy_price == Decimal("25")


def test_ah_high_precision_decimal_arithmetic(orders):
    rows = [tx(units="0.1", nav="0.3"), tx(units="0.1", nav="0.3"), tx(units="0.1", nav="0.3")]
    h = only_holding(parse(orders(rows)))
    # Exact: float arithmetic would give 0.30000000000000004 here.
    assert (h.quantity, h.average_buy_price) == (Decimal("0.3"), Decimal("0.3"))
    rows = [tx(units="1234.567", nav="12.3456"), tx(units="0.001", nav="98765.4321"), tx(kind="REDEEM", units="0.567")]
    h = only_holding(parse(orders(rows)))
    assert h.quantity == Decimal("1234.001")
    expected = B3.divide(Decimal("1234.567") * Decimal("12.3456") + Decimal("0.001") * Decimal("98765.4321"),
                         Decimal("1234.568"))
    assert h.average_buy_price == expected
    assert type(h.quantity) is Decimal and type(h.average_buy_price) is Decimal


def test_repeating_average_keeps_28_significant_digits(orders):
    h = only_holding(parse(orders([tx(units="3", nav="1"), tx(units="3", nav="2"), tx(units="3", nav="2")])))
    assert h.average_buy_price == Decimal("1.666666666666666666666666667")


# --- S, T, U: numeric formats --------------------------------------------------------------------

@pytest.mark.parametrize(
    ("units", "nav", "quantity", "price"),
    [
        ("12.75", "33.5", "12.75", "33.5"),
        (" 12.75 ", " 33.5 ", "12.75", "33.5"),
        ("1,234.567", "1,050.25", "1234.567", "1050.25"),
        ("12,34,567", "10", "1234567", "10"),
        (12.75, 33.5, "12.75", "33.5"),
        (8, 41, "8", "41"),
    ],
    ids=["text", "text-with-spaces", "thousands-separators", "indian-grouping", "float-cells", "integer-cells"],
)
def test_t_u_units_and_nav_as_text_or_numeric_cells(orders, units, nav, quantity, price):
    h = only_holding(parse(orders([tx(units=units, nav=nav)])))
    assert (h.quantity, h.average_buy_price) == (Decimal(quantity), Decimal(price))


@pytest.mark.parametrize("amount", ["2,750", "12,34,567.50", "999", 250.0, "0"])
def test_s_amount_formats_accepted(orders, amount):
    assert parse(orders([tx(amount=amount)])).errors == []


# --- J-R, V: row-level validation ------------------------------------------------------------------

@pytest.mark.parametrize(
    ("changes", "field", "message"),
    [
        ({"units": "0"}, "Units", "Units must be greater than zero"),  # J
        ({"units": "-5"}, "Units", "Units must be greater than zero"),  # K
        ({"kind": "REDEEM", "units": "0"}, "Units", "Units must be greater than zero"),  # L
        ({"kind": "REDEEM", "units": "-5"}, "Units", "Units must be greater than zero"),  # M
        ({"scheme": None}, "Scheme Name", "Scheme Name is missing"),  # N
        ({"scheme": "   "}, "Scheme Name", "Scheme Name is missing"),
        ({"kind": None}, "Transaction Type", "Transaction Type is missing"),  # O
        ({"kind": "  "}, "Transaction Type", "Transaction Type is missing"),
        ({"kind": "SWITCH IN"}, "Transaction Type", UNSUPPORTED_TYPE),  # P
        ({"kind": "DIVIDEND"}, "Transaction Type", UNSUPPORTED_TYPE),
        ({"kind": "purchase"}, "Transaction Type", UNSUPPORTED_TYPE),
        ({"kind": "PURCHASED"}, "Transaction Type", UNSUPPORTED_TYPE),
        ({"units": "abc"}, "Units", "Units is not a number"),  # Q
        ({"units": "1,2"}, "Units", "Units is not a number"),
        ({"units": "1e3"}, "Units", "Units is not a number"),
        ({"units": None}, "Units", "Units is missing or not a number"),
        ({"nav": "N/A"}, "NAV", "NAV is not a number"),  # R
        ({"nav": "-1"}, "NAV", "NAV must not be negative"),
        ({"nav": None}, "NAV", "NAV is missing or not a number"),
        ({"when": "2026-01-15"}, "Date", "Date is not a valid DD Mon YYYY date"),  # V
        ({"when": "15/01/2026"}, "Date", "Date is not a valid DD Mon YYYY date"),
        ({"when": "31 Feb 2026"}, "Date", "Date is not a valid DD Mon YYYY date"),
        ({"when": "5 Jan 2026"}, "Date", "Date is not a valid DD Mon YYYY date"),
        ({"when": "15 JAN 2026"}, "Date", "Date is not a valid DD Mon YYYY date"),
        ({"when": "15 January 2026"}, "Date", "Date is not a valid DD Mon YYYY date"),
        ({"when": None}, "Date", "Date is missing"),
    ],
)
def test_invalid_row_is_reported_and_not_aggregated(orders, changes, field, message):
    # The bad row is the only ALPHA row (or has no scheme at all); BETA is unaffected.
    rows = [tx(scheme=BETA, units="4", nav="25"), tx(**changes), tx(scheme=BETA, units="6", nav="50")]
    result = parse(orders(rows))
    assert [(e.row_number, e.field, e.message) for e in result.errors] == [(HEADER_ROW + 2, field, message)]
    assert result.failed_row_count == 1
    assert [(h.instrument_name, h.quantity, h.average_buy_price) for h in result.holdings] == [
        (BETA, Decimal("10"), Decimal("40"))
    ]


def test_zero_nav_is_accepted(orders):
    h = only_holding(parse(orders([tx(units="10", nav="0"), tx(units="10", nav="20")])))
    assert h.average_buy_price == Decimal("10")


def test_transaction_type_surrounding_whitespace_is_harmless(orders):
    h = only_holding(parse(orders([tx(kind=" PURCHASE ", units="10"), tx(kind="REDEEM\t", units="4")])))
    assert h.quantity == Decimal("6")


def test_every_invalid_field_in_a_row_is_reported(orders):
    bad = tx(scheme=None, kind="SWITCH", units="-1", nav="x", amount="y", when="soon")
    result = parse(orders([bad]))
    assert result.holdings == []
    assert [e.field for e in result.errors] == [h for h in HEADER if h != "Amount"]  # Amount is never read
    assert result.failed_row_count == 1


# --- Scheme-level invalidation ----------------------------------------------------------------------

def test_invalid_purchase_row_excludes_whole_scheme(orders):
    rows = [tx(units="10", nav="20"), tx(units="10", nav="bad"), tx(units="30", nav="40"), tx(scheme=BETA, units="5")]
    result = parse(orders(rows))
    assert [(e.row_number, e.field) for e in result.errors] == [(HEADER_ROW + 2, "NAV")]
    assert [h.instrument_name for h in result.holdings] == [BETA]  # no partial ALPHA from its two valid rows


def test_invalid_redeem_row_excludes_whole_scheme(orders):
    # Aggregating the valid rows alone would overstate ALPHA as 40 units.
    rows = [tx(units="40", nav="20"), tx(kind="REDEEM", units="oops"), tx(scheme=BETA, units="5")]
    result = parse(orders(rows))
    assert [(e.row_number, e.field) for e in result.errors] == [(HEADER_ROW + 2, "Units")]
    assert [h.instrument_name for h in result.holdings] == [BETA]


def test_invalid_scheme_does_not_affect_valid_scheme(orders):
    rows = [tx(scheme=BETA, units="4", nav="10"), tx(kind="SWITCH IN"), tx(units="10", nav="20"),
            tx(scheme=BETA, kind="REDEEM", units="1", nav="99"), tx(scheme=BETA, units="6", nav="15")]
    result = parse(orders(rows))
    assert [e.field for e in result.errors] == ["Transaction Type"]
    assert [(h.instrument_name, h.quantity, h.average_buy_price) for h in result.holdings] == [
        (BETA, Decimal("9"), Decimal("13")),  # (40 + 90) / 10; REDEEM NAV unused
    ]


def test_multiple_invalid_rows_in_one_scheme_all_reported(orders):
    rows = [tx(units="-1"), tx(units="10"), tx(nav="x", when="soon"), tx(kind="REDEEM", units="0"), tx(scheme=BETA)]
    result = parse(orders(rows))
    assert [(e.row_number, e.field) for e in result.errors] == [
        (HEADER_ROW + 1, "Units"), (HEADER_ROW + 3, "NAV"), (HEADER_ROW + 3, "Date"), (HEADER_ROW + 4, "Units"),
    ]
    assert result.failed_row_count == 3
    assert [h.instrument_name for h in result.holdings] == [BETA]


def test_scheme_without_errors_is_calculated_normally_next_to_invalid_scheme(orders):
    clean = [tx(scheme=BETA, units="100", nav="10"), tx(scheme=BETA, units="100", nav="20"),
             tx(scheme=BETA, kind="REDEEM", units="50", nav="100")]
    alone = only_holding(parse(orders(clean)))
    with_bad_neighbour = only_holding(parse(orders([tx(units="bad"), *clean, tx(nav="bad")])))
    assert astuple(alone) == astuple(with_bad_neighbour)
    assert (alone.quantity, alone.average_buy_price) == (Decimal("150"), Decimal("15"))


def test_invalid_row_matches_scheme_after_whitespace_collapse(orders):
    rows = [tx(units="10"), tx(scheme=f"  {ALPHA.replace(' ', '  ')} ", nav="bad"), tx(scheme=BETA)]
    assert [h.instrument_name for h in parse(orders(rows)).holdings] == [BETA]


def test_row_without_scheme_name_is_reported_but_cannot_exclude_a_scheme(orders):
    """A row with no Scheme Name cannot be attributed to any scheme: it is reported, nothing is excluded."""
    result = parse(orders([tx(units="10", nav="20"), tx(scheme=None), tx(scheme=BETA, units="5", nav="8")]))
    assert [(e.row_number, e.field) for e in result.errors] == [(HEADER_ROW + 2, "Scheme Name")]
    assert [h.instrument_name for h in result.holdings] == [ALPHA, BETA]


# --- Amount is never used -----------------------------------------------------------------------------

@pytest.mark.parametrize("amount", ["abc", "-10", None, "   ", "1,2", "N/A"])
def test_malformed_amount_does_not_invalidate_a_valid_transaction(orders, amount):
    rows = [tx(units="100", nav="10", amount=amount), tx(units="100", nav="20"),
            tx(kind="REDEEM", units="50", nav="100", amount=amount)]
    result = parse(orders(rows))
    assert result.errors == []
    assert (only_holding(result).quantity, only_holding(result).average_buy_price) == (Decimal("150"), Decimal("15"))


def test_error_messages_never_echo_cell_values(orders):
    secret = "ACCOUNT-SECRET-99"
    result = parse(orders([{h: secret for h in HEADER}]))
    assert result.errors and all(secret not in e.message for e in result.errors)


def test_personal_details_never_reach_the_output(orders):
    result = parse(orders([tx(), tx(units="bad")]))
    rendered = repr(result.holdings) + repr(result.errors)
    for personal in (HOLDER, MOBILE, PAN):
        assert personal not in rendered


# --- W, X, Y, Z: header detection and layout -------------------------------------------------------

@pytest.mark.parametrize("missing", HEADER)
def test_w_missing_required_header_rejects_file(orders, missing):
    path = orders([tx()], header=[h for h in HEADER if h != missing])
    assert adapter.can_parse(path) is False
    with pytest.raises(StatementFormatError, match="not a Groww mutual fund order history"):
        parse(path)


def test_x_header_found_at_another_row(orders):
    extra = [["Some note"], [], ["Another note"], [], []]
    result = parse(orders([tx(), tx(scheme=BETA, units="bad")], extra_rows_before_header=extra))
    assert len(result.holdings) == 1
    assert [e.row_number for e in result.errors] == [HEADER_ROW + len(extra) + 2]


@pytest.mark.parametrize("column_offset", [1, 4])
def test_y_header_found_in_another_column(write_xlsx, column_offset):
    path = write_xlsx({"Transactions": statement_rows([tx(units="7", nav="3")])}, column_offset=column_offset)
    h = only_holding(parse(path))
    assert (h.quantity, h.average_buy_price) == (Decimal("7"), Decimal("3"))


def test_z_extra_and_reordered_columns(orders):
    header = ["Folio", "Date", "NAV", "Remarks", "Units", "Scheme Name", "Amount", "Transaction Type", "Status"]
    rows = [{**tx(units="10", nav="20"), "Folio": "F-1", "Remarks": "x", "Status": "OK"},
            {**tx(kind="REDEEM", units="4", nav="90"), "Folio": "F-1", "Remarks": "y", "Status": "OK"}]
    h = only_holding(parse(orders(rows, header=header)))
    assert (h.quantity, h.average_buy_price) == (Decimal("6"), Decimal("20"))


# --- Blank rows and empty tables -------------------------------------------------------------------

def test_blank_and_whitespace_rows_skipped(orders):
    blank = {h: None for h in HEADER}
    spaces = {h: "   " for h in HEADER}
    result = parse(orders([blank, tx(), spaces, blank, tx(units="5"), blank]))
    assert (result.errors, only_holding(result).quantity) == ([], Decimal("15"))


def test_empty_transaction_table(orders):
    result = parse(orders([]))
    assert (result.holdings, result.errors, result.statement_date) == ([], [], date(2026, 6, 30))
