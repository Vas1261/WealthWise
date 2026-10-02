"""Groww "Mutual Funds Order History" adapter (spec 3.3) — the one transaction-aggregating adapter.

Groww exports no mutual fund holdings snapshot, only order history, so this adapter (and only
this adapter; equities are never reconstructed from trades) aggregates transactions per scheme
using the locked B3 rule:

    net units          = SUM(PURCHASE units) - SUM(REDEEM units)
    average_buy_price  = SUM(PURCHASE units x PURCHASE NAV) / SUM(PURCHASE units)

REDEEM rows reduce units but never change the average; their NAV is not used, and Amount is
not read at all. Schemes with net units <= 0 are excluded. Invalid rows are reported as row
errors, and a scheme with any invalid row is excluded entirely rather than partially aggregated;
other schemes are unaffected. A row whose Scheme Name itself is missing cannot be attributed to a
scheme, so it is reported but excludes nothing.

Recognition is by content: a sheet qualifies when it has exactly one
"TRANSACTIONS FROM Mon DD YYYY TO Mon DD YYYY" title and, below it, the full transaction header.
The statement date is the title's TO date. The personal details block (name, mobile, PAN) is
never read.
"""

import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_EVEN, Context, Decimal
from typing import TypeVar

from openpyxl.workbook.workbook import Workbook

from app.adapters.base import BrokerAdapter, StatementFormatError
from app.adapters.xlsx import Row, find_header, is_blank_row, load_workbook, normalize_text
from app.domain.holdings import (
    ASSET_CLASS_MUTUAL_FUND,
    InvalidValue,
    NormalizedHolding,
    ParseResult,
    RowError,
    parse_non_negative_decimal,
    parse_positive_decimal,
    parse_required_text,
)

SOURCE_BROKER = "groww"
NATIVE_CURRENCY = "INR"

PREFERRED_SHEET_NAME = "Transactions"  # spec 3.3; only a tie-breaker, never required
STATEMENT_TITLE = re.compile(r"TRANSACTIONS FROM ([A-Z][a-z]{2} \d{2} \d{4}) TO ([A-Z][a-z]{2} \d{2} \d{4})")
TITLE_DATE_FORMAT = "%b %d %Y"  # e.g. "Sep 03 2026"
TRANSACTION_DATE = re.compile(r"\d{2} [A-Z][a-z]{2} \d{4}")
TRANSACTION_DATE_FORMAT = "%d %b %Y"  # e.g. "15 Jan 2026"

SCHEME_NAME = "Scheme Name"
TRANSACTION_TYPE = "Transaction Type"
UNITS = "Units"
NAV = "NAV"
AMOUNT = "Amount"
DATE = "Date"
REQUIRED_HEADERS = (SCHEME_NAME, TRANSACTION_TYPE, UNITS, NAV, AMOUNT, DATE)

PURCHASE = "PURCHASE"
REDEEM = "REDEEM"

# Explicit context so results never depend on the process-wide decimal settings.
DECIMAL_CONTEXT = Context(prec=28, rounding=ROUND_HALF_EVEN)

NOT_A_STATEMENT = (
    "The file is not a Groww mutual fund order history: no sheet has a "
    "'TRANSACTIONS FROM Mon DD YYYY TO Mon DD YYYY' title and a header with " + ", ".join(REQUIRED_HEADERS)
)

T = TypeVar("T")


@dataclass(frozen=True)
class _StatementSheet:
    name: str
    rows: Sequence[Row]
    statement_date: date
    header_index: int
    columns: dict[str, int]


@dataclass(frozen=True)
class _Transaction:
    scheme_name: str
    transaction_type: str
    units: Decimal
    nav: Decimal


@dataclass
class _SchemeTotals:
    purchase_units: Decimal = field(default_factory=Decimal)
    purchase_cost: Decimal = field(default_factory=Decimal)  # SUM(PURCHASE units x PURCHASE NAV)
    redeemed_units: Decimal = field(default_factory=Decimal)


class GrowwMutualFundAdapter(BrokerAdapter):
    def can_parse(self, file_path: str) -> bool:
        workbook = load_workbook(file_path)
        if workbook is None:
            return False
        try:
            _select_statement_sheet(workbook)
        except StatementFormatError:
            return False
        return True

    def parse(self, file_path: str, import_batch_id: str) -> ParseResult:
        workbook = load_workbook(file_path)
        if workbook is None:
            raise StatementFormatError("The file is not a readable .xlsx workbook")
        sheet = _select_statement_sheet(workbook)

        result = ParseResult(statement_date=sheet.statement_date)
        transactions: list[_Transaction] = []
        schemes_with_errors: set[str] = set()
        # Data starts below the header; +1 converts a 0-based index to the Excel row number.
        for index in range(sheet.header_index + 1, len(sheet.rows)):
            row = sheet.rows[index]
            if is_blank_row(row):
                continue
            scheme_name, transaction, errors = _parse_row(row, sheet.columns, index + 1)
            if transaction is not None:
                transactions.append(transaction)
            elif scheme_name is not None:
                schemes_with_errors.add(scheme_name)
            result.errors.extend(errors)
        result.holdings.extend(_aggregate(transactions, schemes_with_errors, import_batch_id))
        return result


def _select_statement_sheet(workbook: Workbook) -> _StatementSheet:
    """Each sheet qualifies on its own; a single qualifying sheet wins, the 'Transactions' name only breaks ties."""
    qualifying = [
        sheet
        for worksheet in workbook.worksheets
        if (sheet := _qualify(worksheet.title, list(worksheet.iter_rows(values_only=True)))) is not None
    ]
    if len(qualifying) > 1:
        qualifying = [sheet for sheet in qualifying if sheet.name == PREFERRED_SHEET_NAME]
        if len(qualifying) != 1:
            raise StatementFormatError("The workbook is ambiguous: more than one qualifying transactions sheet")
    if not qualifying:
        raise StatementFormatError(NOT_A_STATEMENT)
    return qualifying[0]


def _qualify(name: str, rows: Sequence[Row]) -> _StatementSheet | None:
    titles = [
        (row_index, match)
        for row_index, row in enumerate(rows)
        for cell in row
        if (match := STATEMENT_TITLE.fullmatch(normalize_text(cell) or ""))
    ]
    if len(titles) != 1:
        return None
    title_index, match = titles[0]
    period_start, period_end = _title_date(match.group(1)), _title_date(match.group(2))
    if period_start is None or period_end is None or period_start > period_end:
        return None
    header = find_header(rows[title_index + 1 :], REQUIRED_HEADERS)
    if header is None:
        return None
    offset, columns = header
    return _StatementSheet(name, rows, period_end, title_index + 1 + offset, columns)


def _title_date(text: str) -> date | None:
    try:
        return datetime.strptime(text, TITLE_DATE_FORMAT).date()
    except ValueError:
        return None


def _parse_transaction_type(value: object) -> str:
    text = normalize_text(value)
    if text is None:
        raise InvalidValue("is missing")
    if text not in (PURCHASE, REDEEM):
        raise InvalidValue(f"is not supported; only {PURCHASE} and {REDEEM} are imported")
    return text


def _parse_transaction_date(value: object) -> date:
    text = value.strip() if isinstance(value, str) else ""
    if not text:
        raise InvalidValue("is missing")
    if TRANSACTION_DATE.fullmatch(text):
        try:
            return datetime.strptime(text, TRANSACTION_DATE_FORMAT).date()
        except ValueError:
            pass
    raise InvalidValue("is not a valid DD Mon YYYY date")


def _parse_scheme_name(value: object) -> str:
    return " ".join(parse_required_text(value).split())  # group schemes regardless of stray spacing


def _parse_row(
    row: Row, columns: dict[str, int], row_number: int
) -> tuple[str | None, _Transaction | None, list[RowError]]:
    """Validate every field of one transaction row, collecting all failures rather than stopping at the first.

    Returns the row's scheme name whenever it is valid, even if other fields are not, so that an
    invalid row can exclude its whole scheme.
    """
    errors: list[RowError] = []

    def read(header: str, parse: Callable[[object], T]) -> T | None:
        column = columns[header]
        try:
            return parse(row[column] if column < len(row) else None)
        except InvalidValue as error:
            errors.append(RowError(row_number, header, f"{header} {error}"))
            return None

    scheme_name = read(SCHEME_NAME, _parse_scheme_name)
    transaction_type = read(TRANSACTION_TYPE, _parse_transaction_type)
    units = read(UNITS, parse_positive_decimal)
    nav = read(NAV, parse_non_negative_decimal)
    # Date is validated to catch corrupt rows; it does not affect the holding. Amount is never read.
    read(DATE, _parse_transaction_date)

    if errors or scheme_name is None or transaction_type is None or units is None or nav is None:
        return scheme_name, None, errors
    return scheme_name, _Transaction(scheme_name, transaction_type, units, nav), []


def _aggregate(
    transactions: Sequence[_Transaction], excluded_schemes: set[str], import_batch_id: str
) -> list[NormalizedHolding]:
    """Locked B3 rule; schemes keep the order in which they first appear in the file.

    Schemes in excluded_schemes had at least one invalid row and are skipped entirely, so a
    malformed PURCHASE or REDEEM row never yields a misleading partial holding.
    """
    totals: dict[str, _SchemeTotals] = {}
    for transaction in transactions:
        if transaction.scheme_name in excluded_schemes:
            continue
        scheme = totals.setdefault(transaction.scheme_name, _SchemeTotals())
        if transaction.transaction_type == PURCHASE:
            scheme.purchase_units = DECIMAL_CONTEXT.add(scheme.purchase_units, transaction.units)
            scheme.purchase_cost = DECIMAL_CONTEXT.add(
                scheme.purchase_cost, DECIMAL_CONTEXT.multiply(transaction.units, transaction.nav)
            )
        else:  # REDEEM: reduces units only; its NAV is irrelevant to the purchase cost basis
            scheme.redeemed_units = DECIMAL_CONTEXT.add(scheme.redeemed_units, transaction.units)

    holdings = []
    for scheme_name, scheme in totals.items():
        net_units = DECIMAL_CONTEXT.subtract(scheme.purchase_units, scheme.redeemed_units)
        if net_units <= 0:
            continue  # fully redeemed (or more redeemed than purchased in this history): not held
        holdings.append(
            NormalizedHolding(
                instrument_name=scheme_name,
                isin=None,  # not in the export; price lookup will match by scheme (spec 3.3 / 4.5)
                symbol=None,
                asset_class=ASSET_CLASS_MUTUAL_FUND,
                quantity=net_units,
                # net_units > 0 implies purchase_units > 0, so this division is always defined.
                average_buy_price=DECIMAL_CONTEXT.divide(scheme.purchase_cost, scheme.purchase_units),
                native_currency=NATIVE_CURRENCY,
                source_broker=SOURCE_BROKER,
                import_batch_id=import_batch_id,
            )
        )
    return holdings
