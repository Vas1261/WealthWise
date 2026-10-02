"""Groww "Stocks Holdings Statement" adapter (spec 3.1).

A holdings snapshot: quantity and average buy price per stock are read as-is, never
reconstructed from trades. The file's closing price, closing value and P&L columns are
not read: valuation uses live prices (spec 2.3). Personal metadata (name, client code)
is only used to recognise the file and is never returned.
"""

import re
from collections.abc import Callable, Sequence
from datetime import date, datetime
from typing import TypeVar

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.adapters.base import BrokerAdapter, StatementFormatError
from app.adapters.xlsx import Row, find_header, is_blank_row, load_workbook, normalize_text, sheet_texts
from app.domain.holdings import (
    ASSET_CLASS_EQUITY,
    InvalidValue,
    NormalizedHolding,
    ParseResult,
    RowError,
    parse_isin,
    parse_non_negative_decimal,
    parse_positive_decimal,
    parse_required_text,
)

SOURCE_BROKER = "groww"
NATIVE_CURRENCY = "INR"

# Spec 4.2 identification markers.
CLIENT_CODE_LABEL = "Unique Client Code"
STATEMENT_TITLE_PREFIX = "Holdings statement for stocks"
STATEMENT_DATE = re.compile(r"Holdings statement for stocks as on (\d{2}-\d{2}-\d{4})")

STOCK_NAME = "Stock Name"
ISIN = "ISIN"
QUANTITY = "Quantity"
AVERAGE_BUY_PRICE = "Average buy price"
REQUIRED_HEADERS = (STOCK_NAME, ISIN, QUANTITY, AVERAGE_BUY_PRICE)

T = TypeVar("T")


class GrowwEquityAdapter(BrokerAdapter):
    def can_parse(self, file_path: str) -> bool:
        workbook = load_workbook(file_path)
        return workbook is not None and _statement_sheet(workbook) is not None

    def parse(self, file_path: str, import_batch_id: str) -> ParseResult:
        workbook = load_workbook(file_path)
        if workbook is None:
            raise StatementFormatError("The file is not a readable .xlsx workbook")
        sheet = _statement_sheet(workbook)
        if sheet is None:
            raise StatementFormatError("The file is not a Groww stocks holdings statement")

        rows = list(sheet.iter_rows(values_only=True))
        statement_date = _statement_date(rows)
        header = find_header(rows, REQUIRED_HEADERS)
        if header is None:
            raise StatementFormatError(
                "Could not find the holdings table header (" + ", ".join(REQUIRED_HEADERS) + ")"
            )
        header_index, columns = header

        result = ParseResult(statement_date=statement_date)
        # Data starts below the header; +1 converts a 0-based index to the Excel row number.
        for index in range(header_index + 1, len(rows)):
            row = rows[index]
            if is_blank_row(row):
                continue
            holding, errors = _parse_row(row, columns, index + 1, import_batch_id)
            if holding is not None:
                result.holdings.append(holding)
            result.errors.extend(errors)
        return result


def _statement_sheet(workbook: Workbook) -> Worksheet | None:
    """Select the sheet by content (spec 4.2 markers), not by its name."""
    for sheet in workbook.worksheets:
        texts = set(sheet_texts(sheet))
        if CLIENT_CODE_LABEL in texts and any(text.startswith(STATEMENT_TITLE_PREFIX) for text in texts):
            return sheet
    return None


def _statement_date(rows: Sequence[Row]) -> date:
    for row in rows:
        for cell in row:
            match = STATEMENT_DATE.fullmatch(normalize_text(cell) or "")
            if match:
                try:
                    return datetime.strptime(match.group(1), "%d-%m-%Y").date()
                except ValueError:
                    raise StatementFormatError("The statement date is not a valid DD-MM-YYYY date") from None
    raise StatementFormatError("Could not find the 'Holdings statement for stocks as on <date>' line")


def _parse_row(
    row: Row, columns: dict[str, int], row_number: int, import_batch_id: str
) -> tuple[NormalizedHolding | None, list[RowError]]:
    """Validate every required field of one row, collecting all failures rather than stopping at the first."""
    errors: list[RowError] = []

    def read(header: str, parse: Callable[[object], T]) -> T | None:
        column = columns[header]
        try:
            return parse(row[column] if column < len(row) else None)
        except InvalidValue as error:
            errors.append(RowError(row_number, header, f"{header} {error}"))
            return None

    instrument_name = read(STOCK_NAME, parse_required_text)
    isin = read(ISIN, parse_isin)
    quantity = read(QUANTITY, parse_positive_decimal)
    # Average buy price may legitimately be zero (e.g. bonus or gifted shares); only quantity must be positive.
    average_buy_price = read(AVERAGE_BUY_PRICE, parse_non_negative_decimal)

    if instrument_name is None or isin is None or quantity is None or average_buy_price is None:
        return None, errors
    return (
        NormalizedHolding(
            instrument_name=instrument_name,
            isin=isin,
            symbol=None,  # Groww gives no ticker; resolved from ISIN by the Market Service later.
            asset_class=ASSET_CLASS_EQUITY,
            quantity=quantity,
            average_buy_price=average_buy_price,
            native_currency=NATIVE_CURRENCY,
            source_broker=SOURCE_BROKER,
            import_batch_id=import_batch_id,
        ),
        [],
    )
