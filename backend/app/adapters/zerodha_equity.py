"""Zerodha equity holdings adapter (spec 3.2), applying the D1 recognition rule in README.md.

A holdings snapshot: quantity and average price per instrument are read as-is, never
reconstructed from trades. Previous closing price and P&L columns are not read: valuation
uses live prices (spec 2.3). The Client ID label is only used for recognition; its value
is never read.

Sheet selection (D1): a sheet qualifies when it has a `Client ID` label, exactly one title
"Combined|Equity Holdings Statement as on YYYY-MM-DD", and the required header below that
title (Combined also needs `Instrument Type`). Each sheet is judged on its own: one that
fails any requirement simply does not qualify and never affects another sheet. Exactly one
qualifying Combined sheet is used; otherwise exactly one qualifying Equity sheet; two
qualifying sheets of the preferred kind, or none at all, reject the workbook. Only the
selected sheet is read.
"""

import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import date
from typing import TypeVar

from openpyxl.workbook.workbook import Workbook

from app.adapters.base import BrokerAdapter, StatementFormatError
from app.adapters.xlsx import Row, find_header, is_blank_row, load_workbook, normalize_text
from app.domain.holdings import (
    ASSET_CLASS_EQUITY,
    InvalidValue,
    NormalizedHolding,
    ParseResult,
    RowError,
    parse_isin,
    parse_non_negative_decimal,
    parse_positive_decimal,
    parse_required_text,
)

SOURCE_BROKER = "zerodha"
NATIVE_CURRENCY = "INR"

CLIENT_ID_LABEL = "Client ID"
COMBINED = "Combined"
EQUITY = "Equity"
# Matched against whitespace-collapsed cell text, case-sensitively. "Mutual Funds Holdings
# Statement" never matches. The date part is validated separately so a malformed date on a
# supported title is a hard rejection, not a silently ignored line.
STATEMENT_TITLE = re.compile(rf"({COMBINED}|{EQUITY}) Holdings Statement as on(?: (.*))?")
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")

SYMBOL = "Symbol"
ISIN = "ISIN"
QUANTITY_AVAILABLE = "Quantity Available"
AVERAGE_PRICE = "Average Price"
INSTRUMENT_TYPE = "Instrument Type"
REQUIRED_HEADERS = (SYMBOL, ISIN, QUANTITY_AVAILABLE, AVERAGE_PRICE)
EQUITY_INSTRUMENT_TYPE = "-"

NOT_A_STATEMENT = (
    "The file is not a Zerodha equity holdings statement: no sheet has a 'Client ID' label, a "
    "'Combined/Equity Holdings Statement as on YYYY-MM-DD' title and a header with "
    + ", ".join(REQUIRED_HEADERS)
)

T = TypeVar("T")


@dataclass(frozen=True)
class _StatementSheet:
    kind: str  # COMBINED or EQUITY, from the title (never from the sheet name)
    rows: Sequence[Row]
    statement_date: date
    header_index: int
    columns: dict[str, int]


class ZerodhaEquityAdapter(BrokerAdapter):
    def can_parse(self, file_path: str) -> bool:
        workbook = load_workbook(file_path)
        if workbook is None:
            return False
        try:
            _select_statement_sheet(workbook)
        except StatementFormatError:
            return False
        return True

    def parse(self, file_path: str, import_batch_id: str) -> ParseResult:
        workbook = load_workbook(file_path)
        if workbook is None:
            raise StatementFormatError("The file is not a readable .xlsx workbook")
        sheet = _select_statement_sheet(workbook)

        result = ParseResult(statement_date=sheet.statement_date)
        check_instrument_type = sheet.kind == COMBINED
        # Data starts below the header; +1 converts a 0-based index to the Excel row number.
        for index in range(sheet.header_index + 1, len(sheet.rows)):
            row = sheet.rows[index]
            if is_blank_row(row):
                continue
            holding, errors = _parse_row(row, sheet.columns, index + 1, import_batch_id, check_instrument_type)
            if holding is not None:
                result.holdings.append(holding)
            result.errors.extend(errors)
        return result


def _select_statement_sheet(workbook: Workbook) -> _StatementSheet:
    qualifying: dict[str, list[_StatementSheet]] = {COMBINED: [], EQUITY: []}
    reasons: list[str] = []  # why sheets with a supported title did not qualify, for the rejection message
    for worksheet in workbook.worksheets:
        outcome = _qualify(list(worksheet.iter_rows(values_only=True)))
        if isinstance(outcome, _StatementSheet):
            qualifying[outcome.kind].append(outcome)
        elif outcome is not None:
            reasons.append(f"sheet '{worksheet.title}' {outcome}")

    for kind in (COMBINED, EQUITY):
        if len(qualifying[kind]) > 1:
            raise StatementFormatError(f"The workbook is ambiguous: more than one qualifying {kind} holdings sheet")
        if qualifying[kind]:
            return qualifying[kind][0]
    raise StatementFormatError(NOT_A_STATEMENT + "".join(f"; {reason}" for reason in reasons))


def _qualify(rows: Sequence[Row]) -> _StatementSheet | str | None:
    """Judge one sheet against D1 on its own.

    Returns the statement table if the sheet qualifies, a reason if it carries a supported
    title but fails another requirement, or None if it has no supported title at all.
    """
    titles = [
        (row_index, match)
        for row_index, row in enumerate(rows)
        for cell in row
        if (match := STATEMENT_TITLE.fullmatch(normalize_text(cell) or ""))
    ]
    if not titles:
        return None
    if len(titles) > 1:
        return "has more than one holdings statement title"
    title_index, match = titles[0]
    kind = match.group(1)
    statement_date = _iso_date(match.group(2) or "")
    if statement_date is None:
        return "has a statement date that is not a valid YYYY-MM-DD date"
    if not any(normalize_text(cell) == CLIENT_ID_LABEL for row in rows for cell in row):
        return f"has no '{CLIENT_ID_LABEL}' label"
    required = REQUIRED_HEADERS + ((INSTRUMENT_TYPE,) if kind == COMBINED else ())
    header = find_header(rows[title_index + 1 :], required)
    if header is None:
        return "has no header row with " + ", ".join(required) + " below its title"
    offset, columns = header
    return _StatementSheet(kind, rows, statement_date, title_index + 1 + offset, columns)


def _iso_date(text: str) -> date | None:
    if not ISO_DATE.fullmatch(text):
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _parse_row(
    row: Sequence[object],
    columns: dict[str, int],
    row_number: int,
    import_batch_id: str,
    check_instrument_type: bool,
) -> tuple[NormalizedHolding | None, list[RowError]]:
    """Validate every required field of one row, collecting all failures rather than stopping at the first."""
    errors: list[RowError] = []

    def cell(header: str) -> object:
        column = columns[header]
        return row[column] if column < len(row) else None

    def read(header: str, parse: Callable[[object], T]) -> T | None:
        try:
            return parse(cell(header))
        except InvalidValue as error:
            errors.append(RowError(row_number, header, f"{header} {error}"))
            return None

    if check_instrument_type and normalize_text(cell(INSTRUMENT_TYPE)) != EQUITY_INSTRUMENT_TYPE:
        # Only "-" marks an equity row on Combined; anything else (incl. blank) is not imported as equity.
        errors.append(
            RowError(
                row_number,
                INSTRUMENT_TYPE,
                f"{INSTRUMENT_TYPE} is not supported; only equity rows marked '{EQUITY_INSTRUMENT_TYPE}' are imported",
            )
        )
    symbol = read(SYMBOL, parse_required_text)
    isin = read(ISIN, parse_isin)
    quantity = read(QUANTITY_AVAILABLE, parse_positive_decimal)
    # Average price may legitimately be zero (e.g. bonus shares), as for Groww; only quantity must be positive.
    average_buy_price = read(AVERAGE_PRICE, parse_non_negative_decimal)

    if errors or symbol is None or isin is None or quantity is None or average_buy_price is None:
        return None, errors
    return (
        NormalizedHolding(
            instrument_name=symbol,  # Zerodha exports ticker symbols, not full names (spec 3.2)
            isin=isin,
            symbol=symbol,
            asset_class=ASSET_CLASS_EQUITY,
            quantity=quantity,
            average_buy_price=average_buy_price,
            native_currency=NATIVE_CURRENCY,
            source_broker=SOURCE_BROKER,
            import_batch_id=import_batch_id,
        ),
        [],
    )
