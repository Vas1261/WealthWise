"""Shared helpers for reading broker .xlsx exports."""

import warnings
import zipfile
from collections.abc import Iterable, Iterator, Sequence

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

Row = tuple[object, ...]


def load_workbook(file_path: str) -> Workbook | None:
    """Load a workbook fully, or return None if it is not a readable .xlsx file.

    Not read_only mode: some broker exports carry a wrong <dimension> tag, which makes
    openpyxl's streaming reader see an empty 1x1 sheet.
    """
    try:
        with warnings.catch_warnings():
            # Broker exports often omit a default cell style; openpyxl substitutes one, which is harmless.
            warnings.filterwarnings("ignore", message="Workbook contains no default style")
            return openpyxl.load_workbook(file_path, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        return None


def normalize_text(value: object) -> str | None:
    """Cell text with whitespace collapsed; None for non-text or whitespace-only cells."""
    if not isinstance(value, str):
        return None
    text = " ".join(value.split())
    return text or None


def sheet_texts(sheet: Worksheet) -> Iterator[str]:
    for row in sheet.iter_rows(values_only=True):
        for cell in row:
            text = normalize_text(cell)
            if text is not None:
                yield text


def is_blank_row(row: Row) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def find_header(rows: Sequence[Row], required: Iterable[str]) -> tuple[int, dict[str, int]] | None:
    """Locate the first row containing every required header, scanning all rows and columns.

    Requiring the whole set (not one label) avoids matching metadata/summary rows that reuse a
    header word. Returns (0-based row index, {header: 0-based column index}).
    """
    required = set(required)
    for row_index, row in enumerate(rows):
        positions: dict[str, int] = {}
        for column_index, cell in enumerate(row):
            text = normalize_text(cell)
            if text in required and text not in positions:
                positions[text] = column_index
        if len(positions) == len(required):
            return row_index, positions
    return None
